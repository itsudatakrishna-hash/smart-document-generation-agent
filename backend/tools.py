"""Tool implementations for the document generation agent."""

import os
import uuid
import json
from pathlib import Path
from datetime import datetime

from docx import Document
from docx.shared import Pt, Inches
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors

from mock_data import CLIENT_DATA, TEMPLATES, documents_db, signatures_db

OUTPUT_DIR = Path(os.getenv("OUTPUT_DIR", "outputs"))
UPLOAD_DIR = Path(os.getenv("UPLOAD_DIR", "uploads"))
OUTPUT_DIR.mkdir(exist_ok=True)
UPLOAD_DIR.mkdir(exist_ok=True)

TOOL_DEFINITIONS = [
    {
        "type": "function",
        "function": {
            "name": "fetch_client_data",
            "description": "Fetch client information from the CRM database including address, contact details, and billing history.",
            "parameters": {
                "type": "object",
                "properties": {
                    "client_name": {
                        "type": "string",
                        "description": "The client's company name (case-insensitive)",
                    },
                },
                "required": ["client_name"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "fill_template",
            "description": "Fill a .docx template with the provided data and save the result. Returns the file path of the filled document.",
            "parameters": {
                "type": "object",
                "properties": {
                    "template_type": {
                        "type": "string",
                        "enum": ["invoice", "nda", "proposal", "report"],
                        "description": "The type of document template to use",
                    },
                    "data": {
                        "type": "object",
                        "description": "Key-value pairs to fill into the template variables",
                    },
                },
                "required": ["template_type", "data"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "convert_to_pdf",
            "description": "Convert a filled .docx file to PDF format. Returns the PDF file path.",
            "parameters": {
                "type": "object",
                "properties": {
                    "file_path": {
                        "type": "string",
                        "description": "Path to the .docx file to convert",
                    },
                },
                "required": ["file_path"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "upload_to_storage",
            "description": "Upload a file to Google Drive or S3 storage. Returns a mock URL.",
            "parameters": {
                "type": "object",
                "properties": {
                    "file_path": {"type": "string", "description": "Local path to the file"},
                    "destination": {
                        "type": "string",
                        "enum": ["google_drive", "s3"],
                        "description": "Storage destination",
                    },
                },
                "required": ["file_path", "destination"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "send_email",
            "description": "Send the document to a recipient via email.",
            "parameters": {
                "type": "object",
                "properties": {
                    "recipient": {"type": "string", "description": "Recipient email address"},
                    "subject": {"type": "string", "description": "Email subject line"},
                    "file_url": {"type": "string", "description": "URL or path to the attached document"},
                },
                "required": ["recipient", "subject", "file_url"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "request_signature",
            "description": "Send the document for e-signature via DocuSign. Returns a tracking ID.",
            "parameters": {
                "type": "object",
                "properties": {
                    "recipient": {"type": "string", "description": "Signer's email address"},
                    "document_url": {"type": "string", "description": "URL or path to the document"},
                },
                "required": ["recipient", "document_url"],
            },
        },
    },
]


def execute_tool(tool_name: str, tool_input: dict) -> dict:
    if tool_name == "fetch_client_data":
        key = tool_input["client_name"].lower().strip()
        # Fuzzy match: check if any key is contained in the query
        client = CLIENT_DATA.get(key)
        if not client:
            for k, v in CLIENT_DATA.items():
                if k in key or key in k:
                    client = v
                    break
        if not client:
            return {"found": False, "error": f"No client found for '{tool_input['client_name']}'"}
        return {"found": True, **client}

    elif tool_name == "fill_template":
        template_type = tool_input["template_type"]
        data = tool_input["data"]
        doc_id = f"doc-{uuid.uuid4().hex[:8]}"
        output_path = OUTPUT_DIR / f"{doc_id}.docx"
        pdf_path = OUTPUT_DIR / f"{doc_id}.pdf"
        _create_docx(template_type, data, output_path)
        # Always convert to PDF immediately — don't rely on the LLM to call convert_to_pdf
        _convert_docx_to_pdf(output_path, pdf_path)
        size_kb = round(pdf_path.stat().st_size / 1024) if pdf_path.exists() else 0
        documents_db.append({
            "id": doc_id,
            "name": f"{template_type.title()} — {data.get('client_name', 'Document')} — {datetime.now().strftime('%Y-%m-%d')}",
            "type": template_type,
            "client": data.get("client_name", "Unknown"),
            "created_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "status": "generated",
            "file_size": f"{size_kb} KB",
            "file_path": str(pdf_path),
            "signature_status": "not_required",
            "drive_url": None,
            "amount": str(data.get("amount")) if data.get("amount") else None,
        })
        return {
            "success": True,
            "file_path": str(output_path),
            "pdf_path": str(pdf_path),
            "doc_id": doc_id,
            "template_used": TEMPLATES.get(template_type, {}).get("name", template_type),
        }

    elif tool_name == "convert_to_pdf":
        file_path = Path(tool_input["file_path"])
        pdf_path = file_path.with_suffix(".pdf")
        _convert_docx_to_pdf(file_path, pdf_path)
        size_kb = round(pdf_path.stat().st_size / 1024) if pdf_path.exists() else 0
        return {
            "success": True,
            "pdf_path": str(pdf_path),
            "file_size": f"{size_kb} KB",
        }

    elif tool_name == "upload_to_storage":
        file_path = Path(tool_input["file_path"])
        destination = tool_input["destination"]
        file_id = file_path.stem
        if destination == "google_drive":
            url = f"https://drive.google.com/file/d/mock_{file_id}/view"
        else:
            url = f"https://s3.amazonaws.com/my-company-docs/{file_path.name}"
        return {
            "success": True,
            "destination": destination,
            "url": url,
            "file_name": file_path.name,
        }

    elif tool_name == "send_email":
        confirmation_id = f"EMAIL-{uuid.uuid4().hex[:8].upper()}"
        return {
            "success": True,
            "confirmation_id": confirmation_id,
            "recipient": tool_input["recipient"],
            "subject": tool_input["subject"],
            "delivered_at": datetime.now().isoformat(timespec="seconds"),
        }

    elif tool_name == "request_signature":
        tracking_id = f"DS-{uuid.uuid4().hex[:5].upper()}"
        sig_record = {
            "id": f"sig-{uuid.uuid4().hex[:8]}",
            "document_url": tool_input["document_url"],
            "recipient": tool_input["recipient"],
            "sent_at": datetime.now().isoformat(timespec="seconds"),
            "viewed_at": None,
            "signed_at": None,
            "status": "pending",
            "tracking_id": tracking_id,
        }
        signatures_db.append(sig_record)
        return {
            "success": True,
            "tracking_id": tracking_id,
            "recipient": tool_input["recipient"],
            "status": "pending",
            "provider": "DocuSign",
        }

    return {"error": f"Unknown tool: {tool_name}"}


def _create_docx(template_type: str, data: dict, output_path: Path) -> None:
    """Generate a real .docx file using python-docx."""
    doc = Document()

    # Title
    title_map = {
        "invoice": "INVOICE",
        "nda": "NON-DISCLOSURE AGREEMENT",
        "proposal": "BUSINESS PROPOSAL",
        "report": "WEEKLY STATUS REPORT",
    }
    title = doc.add_heading(title_map.get(template_type, "DOCUMENT"), 0)
    title.alignment = 1  # center

    doc.add_paragraph("")

    if template_type == "invoice":
        doc.add_paragraph(f"Invoice Number: {data.get('invoice_number', 'INV-0001')}")
        doc.add_paragraph(f"Date: {datetime.now().strftime('%B %d, %Y')}")
        doc.add_paragraph(f"Due Date: {data.get('due_date', 'Net 30')}")
        doc.add_paragraph("")
        doc.add_heading("Bill To:", 2)
        doc.add_paragraph(data.get("client_name", "Client"))
        doc.add_paragraph(data.get("client_address", ""))
        doc.add_paragraph(data.get("client_email", ""))
        doc.add_paragraph("")
        doc.add_heading("Services", 2)
        doc.add_paragraph(data.get("description", "Professional services rendered"))
        doc.add_paragraph("")
        doc.add_paragraph(f"Amount Due: ${data.get('amount', '0.00')}")

    elif template_type == "nda":
        party_a = data.get('party_a_company') or data.get('party_a', 'Party A')
        party_b = data.get('party_b_company') or data.get('party_b', 'Party B')
        doc.add_paragraph(
            f"This Non-Disclosure Agreement ('Agreement') is entered into as of "
            f"{data.get('effective_date', datetime.now().strftime('%B %d, %Y'))} "
            f"by and between {party_a} and {party_b}."
        )
        doc.add_paragraph("")
        doc.add_heading("1. Confidential Information", 2)
        doc.add_paragraph(
            "Each party may disclose to the other party certain confidential and proprietary information. "
            "The receiving party agrees to keep all such information strictly confidential."
        )
        doc.add_heading("2. Term", 2)
        doc.add_paragraph("This Agreement shall remain in effect for two (2) years from the Effective Date.")
        doc.add_heading("3. Governing Law", 2)
        doc.add_paragraph(f"This Agreement shall be governed by the laws of {data.get('jurisdiction', 'the State of New York')}.")
        doc.add_paragraph("")
        doc.add_paragraph("Signature: ________________________   Date: ____________")

    elif template_type == "proposal":
        doc.add_paragraph(f"Prepared for: {data.get('client_name', 'Client')}")
        doc.add_paragraph(f"Date: {datetime.now().strftime('%B %d, %Y')}")
        doc.add_paragraph("")
        doc.add_heading(data.get("project_title", "Project Title"), 2)
        doc.add_heading("Scope of Work", 2)
        doc.add_paragraph(data.get("scope", "To be defined."))
        doc.add_heading("Investment", 2)
        doc.add_paragraph(f"Total: ${data.get('amount', '0.00')}")
        doc.add_heading("Timeline", 2)
        doc.add_paragraph(f"Estimated Delivery: {data.get('delivery_date', 'TBD')}")

    elif template_type == "report":
        doc.add_paragraph(f"Client: {data.get('client_name', 'Client')}")
        doc.add_paragraph(f"Week Ending: {data.get('week_ending', datetime.now().strftime('%Y-%m-%d'))}")
        doc.add_paragraph("")
        doc.add_heading("Completed This Week", 2)
        doc.add_paragraph(data.get("completed_items", "N/A"))
        doc.add_heading("Next Steps", 2)
        doc.add_paragraph(data.get("next_steps", "N/A"))

    doc.save(str(output_path))


def _convert_docx_to_pdf(docx_path: Path, pdf_path: Path) -> None:
    """Generate a PDF using reportlab (reads the .docx content)."""
    try:
        word_doc = Document(str(docx_path))
        paragraphs_text = [p.text for p in word_doc.paragraphs if p.text.strip()]
    except Exception:
        paragraphs_text = ["Document content"]

    doc = SimpleDocTemplate(str(pdf_path), pagesize=letter,
                            leftMargin=inch, rightMargin=inch,
                            topMargin=inch, bottomMargin=inch)
    styles = getSampleStyleSheet()
    story = []

    for i, text in enumerate(paragraphs_text):
        if i == 0:
            story.append(Paragraph(text, styles["Title"]))
        else:
            story.append(Paragraph(text, styles["Normal"]))
        story.append(Spacer(1, 0.15 * inch))

    doc.build(story)
